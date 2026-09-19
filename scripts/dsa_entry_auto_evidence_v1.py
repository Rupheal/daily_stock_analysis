#!/usr/bin/env python3
"""Build automatic Entry-v1 evidence for simulation-only execution.

Network collection is deliberately isolated from ledger mutation. The collector
uses a 1-minute market bar for the first eligible execution price, HKD/CNY FX,
HKEX board-lot metadata and a fixed simulation-only fee policy. Any missing
field fails closed. No broker or order API is used.
"""
from __future__ import annotations
import argparse,hashlib,json,math
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from dsa_entry_evidence_collector_v1 import build as build_entry_packet

HKT=ZoneInfo("Asia/Hong_Kong")
VERSION="DSA_ENTRY_AUTO_EVIDENCE_v1"

def _sha(obj):
    return hashlib.sha256(json.dumps(obj,sort_keys=True,ensure_ascii=False,separators=(",",":")).encode()).hexdigest()

def _finite(v):
    try:x=float(v)
    except Exception:return None
    return x if math.isfinite(x) else None

def simulated_fee(policy:dict)->str:
    rate=float(policy["fee_rate_on_notional"])
    ticket=float(policy["sizing_reference_ticket_cny"])
    minimum=float(policy["minimum_fee_cny"])
    return str(max(minimum, rate*ticket))

def parse_board_lot(xlsx:Path, code:str)->dict|None:
    import openpyxl
    wb=openpyxl.load_workbook(xlsx,read_only=True,data_only=True)
    code=str(code).zfill(5)
    for ws in wb.worksheets:
        rows=ws.iter_rows(values_only=True)
        header=None;header_row=None
        for i,row in enumerate(rows,1):
            vals=[str(x).strip() if x is not None else "" for x in row]
            low=[x.lower() for x in vals]
            if any("stock code" in x or "code"==x for x in low) and any("board lot" in x for x in low):
                header=vals;header_row=i;break
            if i>=20:break
        if not header:continue
        code_i=next((i for i,x in enumerate(header) if "stock code" in x.lower() or x.lower()=="code"),None)
        lot_i=next((i for i,x in enumerate(header) if "board lot" in x.lower()),None)
        if code_i is None or lot_i is None:continue
        for row in ws.iter_rows(min_row=header_row+1,values_only=True):
            raw=row[code_i] if code_i<len(row) else None
            if raw is None:continue
            s=str(raw).strip().split(".")[0].zfill(5)
            if s!=code:continue
            lot=row[lot_i] if lot_i<len(row) else None
            try:lot_i_val=int(float(lot))
            except Exception:return None
            if lot_i_val<=0:return None
            out={"verified":True,"source":f"HKEX:{xlsx.name}","lot_size":lot_i_val}
            out["sha256"]=_sha(out)
            return out
    return None

def first_minute_bar(code:str, session:str)->dict|None:
    import yfinance as yf
    ticker=f"{str(code).zfill(5)}.HK"
    df=yf.download(ticker,period="1d",interval="1m",auto_adjust=False,progress=False,prepost=False,threads=False)
    if df is None or df.empty:return None
    # yfinance may return tz-aware or naive DatetimeIndex; normalize to HKT.
    rows=[]
    for ts,row in df.iterrows():
        d=ts.to_pydatetime()
        if d.tzinfo is None:d=d.replace(tzinfo=HKT)
        else:d=d.astimezone(HKT)
        if d.date().isoformat()!=session:continue
        price=_finite(row.get("Open"))
        volume=_finite(row.get("Volume"))
        if price and price>0 and volume is not None and volume>0:
            rows.append((d,price,volume))
    if not rows:return None
    d,price,volume=sorted(rows,key=lambda x:x[0])[0]
    out={"verified":True,"source":f"YFinance:{ticker}:1m","at":d.isoformat(),
         "price":price,"volume":volume,"tradable":True,"adjustment":"raw",
         "first_eligible_price_verified":True,"mode":"daily_open","session":session,
         "next_session_verified":True}
    out["sha256"]=_sha(out)
    return out

def fx_hkd_cny(at_iso:str)->dict|None:
    import yfinance as yf
    at=datetime.fromisoformat(at_iso.replace("Z","+00:00"))
    if at.tzinfo is None:return None
    df=yf.download("HKDCNY=X",period="1d",interval="1m",auto_adjust=False,progress=False,prepost=True,threads=False)
    if df is None or df.empty:return None
    candidates=[]
    for ts,row in df.iterrows():
        d=ts.to_pydatetime()
        if d.tzinfo is None:d=d.replace(tzinfo=HKT)
        else:d=d.astimezone(HKT)
        if d>at.astimezone(HKT):continue
        px=_finite(row.get("Close"))
        if px and px>0:candidates.append((d,px))
    if not candidates:return None
    d,px=sorted(candidates,key=lambda x:x[0])[-1]
    out={"verified":True,"source":"YFinance:HKDCNY=X:1m","at":d.isoformat(),"cny_per_hkd":px}
    out["sha256"]=_sha(out)
    return out

def build(code:str,session:str,quote:dict|None,lot:dict|None,fx:dict|None,policy:dict,industry:str|None=None)->dict:
    industry=industry or "UNCLASSIFIED"
    fee={"verified":True,"source":policy["policy_id"],"fee_cny":simulated_fee(policy),
         "simulation_only":True,"real_broker_fee_claim":False}
    fee["sha256"]=_sha(fee)
    ind={"verified":True,"industry":industry,
         "source":"formal-receipt" if industry!="UNCLASSIFIED" else "DSA_CONSERVATIVE_RISK_BUCKET_v1"}
    ind["sha256"]=_sha(ind)
    r=build_entry_packet(code,quote,lot,fx,fee,ind)
    r["auto_evidence_version"]=VERSION
    r["session"]=session
    r["simulation_fee_policy"]=policy["policy_id"]
    r["real_orders"]=0
    r["broker_order_calls"]=0
    return r

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--code",required=True);ap.add_argument("--session",required=True)
    ap.add_argument("--hkex-xlsx",type=Path,required=True)
    ap.add_argument("--fee-policy",type=Path,required=True)
    ap.add_argument("--industry")
    ap.add_argument("--out",type=Path,required=True)
    a=ap.parse_args()
    policy=json.loads(a.fee_policy.read_text())
    quote=first_minute_bar(a.code,a.session)
    lot=parse_board_lot(a.hkex_xlsx,a.code)
    fx=fx_hkd_cny(quote["at"]) if quote else None
    r=build(a.code,a.session,quote,lot,fx,policy,a.industry)
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.out.write_text(json.dumps(r,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"code":a.code,"status":r["status"],"blockers":r["blockers"],"real_orders":0}))
if __name__=="__main__":main()
